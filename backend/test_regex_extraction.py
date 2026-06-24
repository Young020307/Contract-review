"""Verify regex-based extract_fillable_values handles variable-length content.

Character positions verified manually:
- "乙方名称：_____，计费单位：___，单价：_____元"
  乙0 方1 名2 称3 ：4 _5..9 ，10 计11 费12 单13 位14 ：15 _16..18 ，
 19 单20 价21 ：22 _23..27 元28
  fillable: [5,10] [16,19] [23,28]

- "甲方：_____，乙方：_____"
  甲0 方1 ：2 _3..7 ，8 乙9 方10 ：11 _12..16
  fillable: [3,8] [12,17]

- "费用(含税)：_____元/人*天"
  费0 用1 (2 含3 税4 )5 ：6 _7..11 元12 /13 人14 *15 天16
  fillable: [7,12]
"""
import os, sys, tempfile

sys.path.insert(0, os.path.dirname(__file__))

from docx import Document
from services.parser import DocxParser


def make_docx(paragraphs: list[str], filepath: str):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(filepath)


def test_longer_than_placeholder():
    """Values longer than underscore placeholders → full extraction."""
    print("Test 1: longer-than-placeholder values")
    with tempfile.TemporaryDirectory() as tmp:
        tpl = os.path.join(tmp, "tpl.docx")
        doc = os.path.join(tmp, "doc.docx")
        make_docx(["乙方名称：_____，计费单位：___，单价：_____元"], tpl)
        make_docx(["乙方名称：深圳市某某科技有限公司，计费单位：件，单价：5000元"], doc)

        annotations = [
            {"paragraph_index": 0, "start_char": 5, "end_char": 10, "zone_type": "fillable"},
            {"paragraph_index": 0, "start_char": 16, "end_char": 19, "zone_type": "fillable"},
            {"paragraph_index": 0, "start_char": 23, "end_char": 28, "zone_type": "fillable"},
        ]
        values = DocxParser.extract_fillable_values(tpl, doc, annotations)

        assert values["0_5"]["value"] == "深圳市某某科技有限公司", f"FAIL: '{values['0_5']}'"
        assert values["0_16"]["value"] == "件", f"FAIL: '{values['0_16']}'"
        assert values["0_23"]["value"] == "5000", f"FAIL: '{values['0_23']}'"
        # Verify document positions are correct (account for upstream length changes)
        assert values["0_5"]["doc_start"] == 5
        assert values["0_5"]["doc_end"] == 16   # 5 + len("深圳市某某科技有限公司")=11 → 16
        assert values["0_16"]["doc_start"] == 22  # shifted by longer first value
        assert values["0_16"]["doc_end"] == 23
        assert values["0_23"]["doc_start"] == 27
        assert values["0_23"]["doc_end"] == 31
        print("  PASS")


def test_shorter_than_placeholder():
    """Values shorter than placeholder → no extra chars."""
    print("Test 2: shorter-than-placeholder values")
    with tempfile.TemporaryDirectory() as tmp:
        tpl = os.path.join(tmp, "tpl.docx")
        doc = os.path.join(tmp, "doc.docx")
        make_docx(["乙方名称：_____", "金额：__________元"], tpl)
        make_docx(["乙方名称：张三", "金额：100元"], doc)

        annotations = [
            {"paragraph_index": 0, "start_char": 5, "end_char": 10, "zone_type": "fillable"},
            {"paragraph_index": 1, "start_char": 3, "end_char": 13, "zone_type": "fillable"},
        ]
        values = DocxParser.extract_fillable_values(tpl, doc, annotations)

        assert values["0_5"]["value"] == "张三", f"FAIL: '{values['0_5']}'"
        assert values["1_3"]["value"] == "100", f"FAIL: '{values['1_3']}'"
        assert values["0_5"]["doc_start"] == 5
        assert values["0_5"]["doc_end"] == 7   # 5 + len("张三")=2 → 7
        assert values["1_3"]["doc_start"] == 3
        assert values["1_3"]["doc_end"] == 6   # 3 + len("100")=3 → 6
        print("  PASS")


def test_empty_value():
    """User leaves field blank → empty string."""
    print("Test 3: empty (unfilled) values")
    with tempfile.TemporaryDirectory() as tmp:
        tpl = os.path.join(tmp, "tpl.docx")
        doc = os.path.join(tmp, "doc.docx")
        make_docx(["甲方：_____，乙方：_____"], tpl)
        make_docx(["甲方：，乙方：李四"], doc)  # first field empty

        annotations = [
            {"paragraph_index": 0, "start_char": 3, "end_char": 8, "zone_type": "fillable"},
            {"paragraph_index": 0, "start_char": 12, "end_char": 17, "zone_type": "fillable"},
        ]
        values = DocxParser.extract_fillable_values(tpl, doc, annotations)

        assert values["0_3"]["value"] == "", f"FAIL: got '{values['0_3']}'"
        assert values["0_12"]["value"] == "李四", f"FAIL: got '{values['0_12']}'"
        # Empty field: start == end (regex captured zero-length between anchors)
        assert values["0_3"]["doc_start"] == values["0_3"]["doc_end"]
        print("  PASS")


def test_single_zone_whole_paragraph():
    """Single fillable zone spanning almost the whole paragraph."""
    print("Test 4: single zone near-whole paragraph")
    with tempfile.TemporaryDirectory() as tmp:
        tpl = os.path.join(tmp, "tpl.docx")
        doc = os.path.join(tmp, "doc.docx")
        make_docx(["服务项目：", "_____"], tpl)
        make_docx(["服务项目：", "信息系统集成实施服务"], doc)

        annotations = [
            {"paragraph_index": 0, "start_char": 0, "end_char": 5, "zone_type": "fillable"},
            {"paragraph_index": 1, "start_char": 0, "end_char": 5, "zone_type": "fillable"},
        ]
        values = DocxParser.extract_fillable_values(tpl, doc, annotations)

        assert values["0_0"]["value"] == "服务项目：", f"FAIL: '{values['0_0']}'"
        assert values["1_0"]["value"] == "信息系统集成实施服务", f"FAIL: '{values['1_0']}'"
        assert values["1_0"]["doc_start"] == 0
        assert values["1_0"]["doc_end"] == 10  # len("信息系统集成实施服务") = 10
        print("  PASS")


def test_regex_special_chars_in_fixed_text():
    """Fixed text with regex-special chars → properly escaped by re.escape()."""
    print("Test 5: regex-special characters in fixed text")
    with tempfile.TemporaryDirectory() as tmp:
        tpl = os.path.join(tmp, "tpl.docx")
        doc = os.path.join(tmp, "doc.docx")
        make_docx(["费用(含税)：_____元/人*天"], tpl)
        make_docx(["费用(含税)：5000元/人*天"], doc)

        annotations = [
            {"paragraph_index": 0, "start_char": 7, "end_char": 12, "zone_type": "fillable"},
        ]
        values = DocxParser.extract_fillable_values(tpl, doc, annotations)

        assert values["0_7"]["value"] == "5000", f"FAIL: '{values['0_7']}'"
        assert values["0_7"]["doc_start"] == 7
        assert values["0_7"]["doc_end"] == 11  # 7+4 = 11
        print("  PASS")


def test_fallback_on_mismatch():
    """Fixed text differs → regex won't match → fallback to char-slice."""
    print("Test 6: regex mismatch triggers fallback")
    with tempfile.TemporaryDirectory() as tmp:
        tpl = os.path.join(tmp, "tpl.docx")
        doc = os.path.join(tmp, "doc.docx")
        # "方" changed to "万" — regex anchors won't match
        make_docx(["甲方：_____公司"], tpl)
        make_docx(["甲万：某某公司"], doc)

        annotations = [
            {"paragraph_index": 0, "start_char": 3, "end_char": 8, "zone_type": "fillable"},
        ]
        values = DocxParser.extract_fillable_values(tpl, doc, annotations)

        assert values["0_3"]["value"] != "", f"FAIL: fallback should still extract something"
        print(f"  PASS (fallback: '{values['0_3']}')")


def test_no_fillable_annotations():
    """Zero fillable annotations → empty dict."""
    print("Test 7: no fillable annotations")
    with tempfile.TemporaryDirectory() as tmp:
        tpl = os.path.join(tmp, "tpl.docx")
        doc = os.path.join(tmp, "doc.docx")
        make_docx(["hello world"], tpl)
        make_docx(["hello world"], doc)

        values = DocxParser.extract_fillable_values(tpl, doc, [])
        assert values == {}, f"FAIL: {values}"

        values2 = DocxParser.extract_fillable_values(
            tpl, doc,
            [{"paragraph_index": 0, "start_char": 0, "end_char": 5, "zone_type": "fixed"}]
        )
        assert values2 == {}, f"FAIL: {values2}"
        print("  PASS")


def test_paragraph_count_mismatch():
    """Document has fewer paragraphs than template → skip gracefully."""
    print("Test 8: paragraph index out of range")
    with tempfile.TemporaryDirectory() as tmp:
        tpl = os.path.join(tmp, "tpl.docx")
        doc = os.path.join(tmp, "doc.docx")
        make_docx(["para0", "para1", "para2"], tpl)
        make_docx(["para0"], doc)

        annotations = [
            {"paragraph_index": 0, "start_char": 0, "end_char": 5, "zone_type": "fillable"},
            {"paragraph_index": 2, "start_char": 0, "end_char": 5, "zone_type": "fillable"},
        ]
        values = DocxParser.extract_fillable_values(tpl, doc, annotations)
        assert "0_0" in values, f"FAIL: para 0 should be extracted"
        assert "2_0" not in values, f"FAIL: para 2 should be skipped"
        print("  PASS")


if __name__ == "__main__":
    all_pass = True
    tests = [
        test_longer_than_placeholder,
        test_shorter_than_placeholder,
        test_empty_value,
        test_single_zone_whole_paragraph,
        test_regex_special_chars_in_fixed_text,
        test_fallback_on_mismatch,
        test_no_fillable_annotations,
        test_paragraph_count_mismatch,
    ]
    for t in tests:
        try:
            t()
        except Exception as e:
            print(f"  FAIL: {e}")
            all_pass = False

    print()
    if all_pass:
        print("ALL TESTS PASSED")
    else:
        print("SOME TESTS FAILED")
        sys.exit(1)
